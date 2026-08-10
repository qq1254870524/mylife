from __future__ import annotations

import csv
import os
from pathlib import Path

from input_loader import _encoding


def _rewrite_csv(path: Path, completed: set[str]) -> int:
    raw = path.read_bytes()
    encoding = _encoding(raw)
    text = raw.decode(encoding)
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(text.splitlines(), dialect))
    if not rows:
        return 0
    kept = [rows[0]]
    removed = 0
    for row in rows[1:]:
        first = str(row[0] if row else "").strip()
        if first and first in completed:
            removed += 1
        else:
            kept.append(row)
    temp = path.with_name(f".{path.name}.mylife_rebuild.tmp")
    with temp.open("w", encoding=encoding, newline="") as handle:
        writer = csv.writer(handle, dialect)
        writer.writerows(kept)
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temp, path)
    return removed


def _rewrite_txt(path: Path, completed: set[str]) -> int:
    raw = path.read_bytes()
    encoding = _encoding(raw)
    lines = raw.decode(encoding).splitlines()
    removed = 0
    kept: list[str] = []
    for line in lines:
        first = line.strip()
        if first and first in completed:
            removed += 1
        else:
            kept.append(line)
    temp = path.with_name(f".{path.name}.mylife_rebuild.tmp")
    with temp.open("w", encoding=encoding, newline="") as handle:
        handle.write("\n".join(kept) + ("\n" if kept else ""))
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temp, path)
    return removed


def _rewrite_xlsx(path: Path, completed: set[str]) -> int:
    import openpyxl

    keep_vba = path.suffix.lower() == ".xlsm"
    workbook = openpyxl.load_workbook(path, keep_vba=keep_vba)
    removed = 0
    try:
        for sheet in workbook.worksheets:
            delete_rows: list[int] = []
            for row_index in range(2, sheet.max_row + 1):
                first = str(sheet.cell(row_index, 1).value or "").strip()
                if first and first in completed:
                    delete_rows.append(row_index)
            for row_index in reversed(delete_rows):
                sheet.delete_rows(row_index, 1)
            removed += len(delete_rows)
        temp = path.with_name(f".{path.stem}.mylife_rebuild{path.suffix}")
        workbook.save(temp)
    finally:
        workbook.close()
    os.replace(temp, path)
    return removed


def remove_completed_rows(path: str | Path, completed_first_values: set[str]) -> int:
    """仅在整批处理结束后调用；通过同目录临时文件一次性原子重建输入。"""

    source = Path(path).resolve()
    completed = {str(value).strip() for value in completed_first_values if str(value).strip()}
    if not completed:
        return 0
    suffix = source.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _rewrite_xlsx(source, completed)
    if suffix == ".csv":
        return _rewrite_csv(source, completed)
    if suffix == ".txt":
        return _rewrite_txt(source, completed)
    raise ValueError(f"不支持重建的文件类型：{suffix}")
